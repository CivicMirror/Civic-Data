#!/usr/bin/env python3
"""Find Texas ISD board pages from website URLs in an Excel workbook.

The input workbook is never modified. Results are appended to a copied workbook.
"""

from __future__ import annotations

import argparse
import json
import os
import random
import shutil
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from openpyxl import load_workbook


RESULT_HEADERS = (
    "Board Page Result URL",
    "Board Page Result Title",
    "Board Page Domain Match",
    "Board Page Search Status",
    "Board Page Search Error",
    "Board Page Search Query (audit only)",
)


@dataclass
class SearchResult:
    url: str = ""
    title: str = ""


class SearchError(RuntimeError):
    pass


def http_json(url: str, *, method: str = "GET", headers: dict[str, str] | None = None,
              body: dict[str, Any] | None = None, attempts: int = 4) -> dict[str, Any]:
    payload = json.dumps(body).encode("utf-8") if body is not None else None
    request_headers = {"User-Agent": "TX-ISD-Board-Page-Finder/1.0"}
    request_headers.update(headers or {})
    for attempt in range(attempts):
        try:
            req = Request(url, data=payload, headers=request_headers, method=method)
            with urlopen(req, timeout=30) as response:
                return json.loads(response.read().decode("utf-8"))
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", "replace")[:500]
            if exc.code not in (429, 500, 502, 503, 504) or attempt == attempts - 1:
                raise SearchError(f"HTTP {exc.code}: {detail}") from exc
        except (URLError, TimeoutError, json.JSONDecodeError) as exc:
            if attempt == attempts - 1:
                raise SearchError(str(exc)) from exc
        time.sleep((2 ** attempt) + random.random())
    raise SearchError("Search request failed")


def search_serper(query: str, api_key: str) -> SearchResult:
    data = http_json(
        "https://google.serper.dev/search",
        method="POST",
        headers={"X-API-KEY": api_key, "Content-Type": "application/json"},
        body={"q": query, "num": 10},
    )
    organic = data.get("organic") or []
    if not organic:
        return SearchResult()
    first = organic[0]
    return SearchResult(url=str(first.get("link", "")), title=str(first.get("title", "")))


def search_google_cse(query: str, api_key: str, cx: str) -> SearchResult:
    from urllib.parse import urlencode

    endpoint = "https://customsearch.googleapis.com/customsearch/v1?" + urlencode(
        {"key": api_key, "cx": cx, "q": query, "num": 1}
    )
    data = http_json(endpoint)
    items = data.get("items") or []
    if not items:
        return SearchResult()
    first = items[0]
    return SearchResult(url=str(first.get("link", "")), title=str(first.get("title", "")))


def hostname(value: str) -> str:
    parsed = urlparse(value if "://" in value else f"https://{value}")
    return (parsed.hostname or "").lower().removeprefix("www.")


def same_domain(source: str, result: str) -> bool:
    source_host, result_host = hostname(source), hostname(result)
    return bool(source_host and result_host and (
        source_host == result_host
        or source_host.endswith("." + result_host)
        or result_host.endswith("." + source_host)
    ))


def find_header_row(ws, expected: str, max_rows: int = 25) -> int:
    for row in range(1, max_rows + 1):
        if str(ws.cell(row, 1).value or "").strip().casefold() == expected.casefold():
            return row
    raise ValueError(f"Could not find '{expected}' in column A")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Find the first Google result for each ISD website.")
    parser.add_argument("workbook", type=Path, help="TX_Municipalities .xlsx file")
    parser.add_argument("-o", "--output", type=Path, help="Output .xlsx path")
    parser.add_argument("--provider", choices=("serper", "google-cse"), default="serper")
    parser.add_argument("--search-term", default="board of trustees page")
    parser.add_argument("--sheet", default="ISD")
    parser.add_argument("--website-column", default="D")
    parser.add_argument("--delay", type=float, default=1.0, help="Seconds between searches")
    parser.add_argument("--limit", type=int, help="Maximum rows to search (useful for testing)")
    parser.add_argument("--start-row", type=int, help="First worksheet row to process")
    parser.add_argument("--overwrite-results", action="store_true")
    parser.add_argument("--dry-run", action="store_true", help="Print queries without calling an API")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.workbook.exists():
        print(f"Input not found: {args.workbook}", file=sys.stderr)
        return 2

    output = args.output or args.workbook.with_name(args.workbook.stem + "_board_pages.xlsx")
    if args.workbook.resolve() == output.resolve():
        print("Output must be different from input.", file=sys.stderr)
        return 2

    if args.provider == "serper":
        api_key = os.getenv("SERPER_API_KEY", "")
        cx = ""
    else:
        api_key = os.getenv("GOOGLE_CSE_API_KEY", "")
        cx = os.getenv("GOOGLE_CSE_ID", "")
    if not args.dry_run and (not api_key or (args.provider == "google-cse" and not cx)):
        needed = "SERPER_API_KEY" if args.provider == "serper" else "GOOGLE_CSE_API_KEY and GOOGLE_CSE_ID"
        print(f"Missing environment variable(s): {needed}", file=sys.stderr)
        return 2

    output.parent.mkdir(parents=True, exist_ok=True)
    if not output.exists():
        shutil.copy2(args.workbook, output)

    wb = load_workbook(output)
    if args.sheet not in wb.sheetnames:
        print(f"Sheet not found: {args.sheet}", file=sys.stderr)
        return 2
    ws = wb[args.sheet]
    header_row = find_header_row(ws, "ISD Name")
    header_map = {str(ws.cell(header_row, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    next_col = ws.max_column + 1
    for label in RESULT_HEADERS:
        if label not in header_map:
            header_map[label] = next_col
            ws.cell(header_row, next_col, label)
            next_col += 1

    url_col = header_map[RESULT_HEADERS[0]]
    title_col = header_map[RESULT_HEADERS[1]]
    match_col = header_map[RESULT_HEADERS[2]]
    status_col = header_map[RESULT_HEADERS[3]]
    error_col = header_map[RESULT_HEADERS[4]]
    query_col = header_map[RESULT_HEADERS[5]]
    first_row = args.start_row or (header_row + 1)
    processed = 0

    for row in range(first_row, ws.max_row + 1):
        website = str(ws[f"{args.website_column}{row}"].value or "").strip()
        if not website:
            continue
        existing_status = str(ws.cell(row, status_col).value or "").strip()
        if existing_status and not args.overwrite_results:
            continue
        if args.limit is not None and processed >= args.limit:
            break

        query = f"{website} {args.search_term}"
        district = str(ws.cell(row, 1).value or "")
        print(f"[{row}/{ws.max_row}] {district}: {query}", flush=True)
        if args.dry_run:
            processed += 1
            continue

        ws.cell(row, query_col, query)
        try:
            result = (search_serper(query, api_key) if args.provider == "serper"
                      else search_google_cse(query, api_key, cx))
            ws.cell(row, url_col, result.url or None)
            ws.cell(row, title_col, result.title or None)
            ws.cell(row, match_col, "Yes" if result.url and same_domain(website, result.url) else "No")
            ws.cell(row, status_col, "found" if result.url else "no_result")
            ws.cell(row, error_col, None)
        except SearchError as exc:
            ws.cell(row, status_col, "error")
            ws.cell(row, error_col, str(exc)[:1000])

        processed += 1
        # Save after every result so an interrupted run can resume safely.
        wb.save(output)
        if args.delay > 0:
            time.sleep(args.delay)

    if not args.dry_run:
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = ws.dimensions
        wb.save(output)
        print(f"Saved {processed} processed rows to {output}")
    else:
        print(f"Dry run complete: {processed} queries")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
